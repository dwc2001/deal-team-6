"""One-time import of the original Deal Team 6 spreadsheet.

Reads "Deal Team 6 - Tools and Resouces.xlsx", cleans it up, and writes:
  data/seed.local.json      the cleaned records (used by local demo mode)
  supabase/seed.local.sql   the same records as SQL to paste into Supabase

Both outputs hold real contact details, so they are gitignored. Run:
  pip install openpyxl
  python scripts/import_excel.py
"""

from __future__ import annotations

import json
import re
import uuid
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "Deal Team 6 - Tools and Resouces.xlsx"
OUT_JSON = ROOT / "data" / "seed.local.json"
OUT_SQL = ROOT / "supabase" / "seed.local.sql"

IMPORTED_BY = "Spreadsheet"
IMPORTED_AT = "2026-10-05T12:00:00Z"

# Every service the sheet uses, and the group it is listed under on the Directory.
GROUPS = {
    "Lending and 1031": [
        "1031 / QI", "1031 STNL/DST Buyside Specialists", "DST",
        "Loan Officer", "Loan Officer (SBA)",
    ],
    "Legal, tax and title": [
        "Accountant", "Attorney", "Cost Segregation", "Land Use Attorney (PG Co.)",
        "Tenant / TOPA Negotiations", "Title Company",
    ],
    "Brokers and agents": ["Business Broker", "Condo Outsales", "Leasing", "Resi Agent"],
    "Due diligence": [
        "Appraisal", "Engineer", "Environmental", "Flood Zone Evaluation / Resolution",
        "Survey", "Valuation / Biz Planning",
    ],
    "Property services": [
        "Cleaning Services", "Contractor", "Insurance Broker", "Inventory Liquidation",
        "Junk Removal", "Landscaping / Snow Removal", "Property Management",
        "Sign and Promo Items",
    ],
    "Everything else": ["Car Service", "Chef / Personal / Events"],
}

CATEGORY_FIXES = {
    "Enviormental": "Environmental",
    "Property Managment": "Property Management",
    "Property Insurance Broker": "Insurance Broker",
    "Landscaping/Snow Removal": "Landscaping / Snow Removal",
}

# Row-specific fixes (who knows whom, who pays referrals, typos in single rows)
# hold real names, so they live in scripts/overrides.local.json, not in git.
OVERRIDES_FILE = ROOT / "scripts" / "overrides.local.json"
OVERRIDES = json.loads(OVERRIDES_FILE.read_text(encoding="utf-8")) if OVERRIDES_FILE.exists() else {}
CONNECTIONS: dict[str, str] = OVERRIDES.get("connections", {})
PAYS_REFERRAL: set[str] = set(OVERRIDES.get("pays_referral", []))
CONTACT_FIXES: dict[str, dict[str, str]] = OVERRIDES.get("contact_fixes", {})
TENANT_IN_ADDRESS: dict[str, list[str]] = OVERRIDES.get("tenant_in_address", {})
USE_IS_TENANT: dict[str, str] = OVERRIDES.get("use_is_tenant", {})

# Notes that only repeated the service name.
NOISE_NOTES = {
    "appraisal", "condo outsales", "loan officer", "insurance broker",
    "chef / personal / events",
}


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = str(value).replace("‬", "").replace("‑", "-").replace(" ", " ")
    s = re.sub(r"\s+", " ", s).strip()
    return "" if s in {"-", "?", "IDK", "who knows"} else s


def phone(value) -> str:
    raw = text(value)
    digits = re.sub(r"\D", "", raw)
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    if len(digits) == 10:
        return f"({digits[:3]}) {digits[3:6]}-{digits[6:]}"
    return raw


def iso(value) -> str | None:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return None


def num(value) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    return None


def percent(value) -> str:
    if isinstance(value, (int, float)):
        pct = value * 100 if value < 1 else value
        return f"{pct:g}%"
    return text(value)


def money_psf(value) -> str:
    if isinstance(value, (int, float)):
        return "None" if value == 0 else f"${value:,.2f}/SF".replace(".00/", "/")
    s = text(value)
    m = re.fullmatch(r"\$?(\d+(?:\.\d+)?)\s*/\s*sf", s, re.I)
    return f"${m.group(1)}/SF" if m else s


def months(value) -> str:
    if isinstance(value, (int, float)):
        if value == 0:
            return "None"
        return "1 month" if value == 1 else f"{value:g} months"
    s = text(value)
    if s == "1 yr":
        return "12 months"
    return s


def years(value) -> tuple[int | None, int | None]:
    s = text(value)
    parts = [int(p) for p in re.findall(r"\d{4}", s)]
    if not parts:
        return None, None
    return parts[0], (parts[1] if len(parts) > 1 else None)


# ----------------------------------------------------------------- contacts
def read_contacts(ws) -> list[dict]:
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        service, name, company, number, email, territory, title, notes = (list(r) + [None] * 8)[:8]
        service = text(service)
        if not service and not text(name):
            continue
        preferred = service.endswith("*")
        service = service.rstrip("*").strip()
        service = CATEGORY_FIXES.get(service, service)
        name = text(name)
        note = text(notes)
        if note.lower() in NOISE_NOTES:
            note = ""
        company_s = text(company)
        title_s = text(title)

        fix = CONTACT_FIXES.get(name, {})
        connection = CONNECTIONS.get(fix.get("name", name), "")
        referral = fix.get("name", name) in PAYS_REFERRAL or name in PAYS_REFERRAL
        name = fix.get("name", name)
        company_s = fix.get("company", company_s)
        note = fix.get("notes", note)

        rows.append({
            "id": str(uuid.uuid5(uuid.NAMESPACE_URL, f"dt6-contact-{service}-{name}")),
            "category": service,
            "name": name,
            "company": company_s,
            "title": title_s,
            "phone": phone(number),
            "phone_alt": "",
            "email": text(email).lower() if text(email) else "",
            "website": "",
            "address": "",
            "territory": text(territory).replace("World Wide", "Worldwide"),
            "notes": note,
            "connection": connection,
            "preferred": preferred,
            "pays_referral": referral,
            "card_front": None,
            "added_by": IMPORTED_BY,
            "updated_by": "",
            "archived": False,
            "created_at": IMPORTED_AT,
            "updated_at": IMPORTED_AT,
        })
    return rows


# -------------------------------------------------------------------- comps


def lease_type(value) -> tuple[str, str]:
    s = text(value)
    if not s:
        return "", ""
    if s.lower().startswith("mod"):
        extra = "Tenant pays utilities." if "utilit" in s.lower() else ""
        return "Modified Gross", extra
    return s, ""


def comp_base(**kw) -> dict:
    base = {
        "id": "", "address": "", "center_name": "", "submarket": "", "city": "",
        "state": "", "county": "", "space_type": "", "use_type": "", "tenant": "",
        "year_built": None, "year_renovated": None, "sign_date": None, "start_date": None,
        "sf": None, "floor": "", "rate_psf": None, "term_months": None, "lease_type": "",
        "deal_type": "", "escalations": "", "options": "", "ti": "", "free_rent": "",
        "notes": "", "added_by": IMPORTED_BY, "updated_by": "", "archived": False,
        "created_at": IMPORTED_AT, "updated_at": IMPORTED_AT,
    }
    base.update(kw)
    return base


def read_commercial(ws) -> list[dict]:
    out = []
    for i, r in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        (_, address, hood, use, built, sign, start, sf, floor, rate, term_m, term_y,
         basis, esc, options, ti, free, tenant) = (list(r) + [None] * 18)[:18]
        address = text(address)
        if not address:
            continue
        tenant = text(tenant)
        if address in TENANT_IN_ADDRESS:
            address, tenant = TENANT_IN_ADDRESS[address]
        address = address.replace("1001Pennsylvania", "1001 Pennsylvania")
        hood = text(hood).replace("MT Vernon", "Mt Vernon")
        use_s = text(use)
        space_type, use_type = ("Retail", "Restaurant") if "Restaurant" in use_s else (use_s, "")
        y_built, y_reno = years(built)

        # The sheet stores formulas; recompute what they said.
        if isinstance(rate, str) and rate.startswith("="):
            rate = 78000 / float(sf)
        if isinstance(term_m, (int, float)):
            term = int(term_m)
        elif isinstance(term_y, (int, float)):
            term = int(round(term_y * 12))
        else:
            term = None

        lt, lt_note = lease_type(basis)
        notes = [lt_note] if lt_note else []
        start_iso = iso(start)
        sign_iso = iso(sign)
        # Three Bethesda rows put the lease expiration in the start date column.
        if start_iso and sign_iso and start_iso[:4] >= "2033":
            notes.append(f"Runs through {start.strftime('%b %d, %Y')} (listed as the start date in the original sheet).")
            start_iso = None
        ti_s = money_psf(ti)
        if ti_s == "% Rent, yr 1":
            notes.append("Percentage rent in year 1.")
            ti_s = ""

        is_md = hood == "Bethesda"
        out.append(comp_base(
            id=str(uuid.uuid5(uuid.NAMESPACE_URL, f"dt6-comp-c{i}-{address}")),
            address=address,
            submarket=hood,
            city="Bethesda" if is_md else "Washington",
            state="MD" if is_md else "DC",
            county="Montgomery" if is_md else "",
            space_type=space_type,
            use_type=use_type,
            tenant=tenant,
            year_built=y_built,
            year_renovated=y_reno,
            sign_date=sign_iso,
            start_date=start_iso,
            sf=num(sf),
            floor=text(floor),
            rate_psf=round(float(rate), 2) if isinstance(rate, (int, float)) else None,
            term_months=term,
            lease_type=lt,
            escalations=percent(esc),
            options=text(options),
            ti=ti_s,
            free_rent=months(free),
            notes=" ".join(notes),
        ))
    return out


def read_retail(ws) -> list[dict]:
    out = []
    for i, r in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        (_, address, center, city, state, county, use, built, sign, start, sf, rate,
         term_m, term_y, ltype, new_renew, esc, ti, free) = (list(r) + [None] * 19)[:19]
        address = text(address)
        if not address:
            continue
        use_s = text(use)
        tenant, use_type = "", use_s
        if use_s in USE_IS_TENANT:
            tenant, use_type = use_s, USE_IS_TENANT[use_s]
        lt, lt_note = lease_type(ltype)
        y_built, y_reno = years(built)
        out.append(comp_base(
            id=str(uuid.uuid5(uuid.NAMESPACE_URL, f"dt6-comp-r{i}-{address}")),
            address=address,
            center_name=text(center),
            submarket=text(city).replace("Skyesville", "Sykesville"),
            city=text(city).replace("Skyesville", "Sykesville"),
            state=text(state),
            county=text(county),
            space_type="Retail",
            use_type=use_type,
            tenant=tenant,
            year_built=y_built,
            year_renovated=y_reno,
            sign_date=iso(sign),
            start_date=iso(start),
            sf=num(sf),
            rate_psf=num(rate),
            term_months=int(term_m) if isinstance(term_m, (int, float)) else None,
            lease_type=lt,
            deal_type=text(new_renew),
            escalations=percent(esc),
            ti=money_psf(ti),
            free_rent=months(free),
            notes=lt_note,
        ))
    return out


# -------------------------------------------------------------------- links
def clean_url(url) -> str:
    s = text(url)
    # Links saved from Acrobat's Chrome extension wrap the real address.
    return re.sub(r"^chrome-extension://[a-z]+/", "", s)


def read_links(grants_ws, sites_ws) -> list[dict]:
    out = []
    order = 0
    for r in grants_ws.iter_rows(min_row=3, values_only=True):
        flag, name, site, brochure, map_url = (list(r) + [None] * 5)[:5]
        name = text(name).rstrip(":").replace("Porgram", "Program")
        if not name:
            continue
        more = []
        if clean_url(brochure):
            more.append({"label": "Brochure", "url": clean_url(brochure)})
        if clean_url(map_url):
            more.append({"label": "Map", "url": clean_url(map_url)})
        order += 1
        out.append({
            "id": str(uuid.uuid5(uuid.NAMESPACE_URL, f"dt6-link-{name}")),
            "section": "Grants and incentives",
            "jurisdiction": "DC",
            "title": name,
            "url": clean_url(site),
            "description": "",
            "status": "Closed" if "closed" in text(flag).lower() else "",
            "more": more,
            "sort": order,
            "added_by": IMPORTED_BY, "updated_by": "", "archived": False,
            "created_at": IMPORTED_AT, "updated_at": IMPORTED_AT,
        })
    descriptions = {
        "Certificate of Occupancy Finder": "DC records search for certificates of occupancy.",
        "Office Conversion Primer": "Greysteel's primer on office conversions (Box).",
        "Closing Cost Estimator": "Shulman Rogers title quote calculator.",
        "Business Day Calculator": "Count business days between two dates.",
    }
    for r in sites_ws.iter_rows(min_row=1, values_only=True):
        name, url = (list(r) + [None] * 2)[:2]
        name = text(name)
        if not name:
            continue
        order += 1
        out.append({
            "id": str(uuid.uuid5(uuid.NAMESPACE_URL, f"dt6-link-{name}")),
            "section": "Tools and websites",
            "jurisdiction": "",
            "title": name,
            "url": clean_url(url),
            "description": descriptions.get(name, ""),
            "status": "",
            "more": [],
            "sort": order,
            "added_by": IMPORTED_BY, "updated_by": "", "archived": False,
            "created_at": IMPORTED_AT, "updated_at": IMPORTED_AT,
        })
    return out


# ---------------------------------------------------------------------- SQL
def sql_value(v) -> str:
    if v is None:
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return repr(v)
    if isinstance(v, (list, dict)):
        return "'" + json.dumps(v).replace("'", "''") + "'::jsonb"
    return "'" + str(v).replace("'", "''") + "'"


def sql_insert(table: str, rows: list[dict]) -> str:
    if not rows:
        return ""
    cols = list(rows[0].keys())
    values = ",\n".join("  (" + ", ".join(sql_value(row[c]) for c in cols) + ")" for row in rows)
    return f"insert into public.{table} ({', '.join(cols)}) values\n{values}\non conflict do nothing;\n"


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, data_only=False)
    contacts = read_contacts(wb["Vendor List"])
    comps = read_commercial(wb["Lease Comps (Commercial)"]) + read_retail(wb["Leas Comps (Retail)"])
    links = read_links(wb["Grant Info"], wb["Websites"])

    group_of = {c: g for g, cats in GROUPS.items() for c in cats}
    used = sorted({c["category"] for c in contacts})
    missing = [c for c in used if c not in group_of]
    if missing:
        raise SystemExit(f"Add these services to GROUPS: {missing}")
    categories = [{"name": c, "grp": group_of[c]} for c in used]

    data = {"categories": categories, "contacts": contacts, "comps": comps, "links": links}
    OUT_JSON.parent.mkdir(exist_ok=True)
    OUT_JSON.write_text(json.dumps(data, indent=2), encoding="utf-8")

    sql = [
        "-- Deal Team 6: records imported from the original spreadsheet.",
        "-- Paste into the Supabase SQL editor after schema.sql. Safe to run twice.",
        "",
        sql_insert("categories", categories),
        sql_insert("contacts", contacts),
        sql_insert("comps", comps),
        sql_insert("links", links),
    ]
    OUT_SQL.write_text("\n".join(sql), encoding="utf-8")
    print(f"{len(contacts)} contacts in {len(categories)} services, {len(comps)} comps, {len(links)} links")
    print(f"wrote {OUT_JSON.relative_to(ROOT)} and {OUT_SQL.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
